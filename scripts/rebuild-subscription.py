"""Editable Word templates + measured PDF maps. Run with bundled Python.
The probe uses identical fixed rows; only the blank answer paragraphs contain markers.
Their PDF positions determine blue overlay areas; final Word/PDF contain no markers.
"""
from pathlib import Path
import json, subprocess, shutil, zipfile, uuid
from lxml import etree
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pdfplumber

ROOT=Path(__file__).resolve().parents[1]
RUNTIME=Path('/Users/ahmaddalao/.cache/codex-runtimes/codex-primary-runtime/dependencies')
RENDER=Path('/Users/ahmaddalao/.codex/plugins/cache/openai-primary-runtime/documents/26.904.11930/skills/documents/render_docx.py')
RULES=json.loads((ROOT/'public/api/subscription/rules.json').read_text())
FONT='Bahij TheSansArabic Plain'
PURPLE='482563';INK='242235';MUTED='656575';BORDER='D8DCE5'
WIDTH=499.3
OPEN_AR='بعد الاطلاع على نشرة الشروط والأحكام للصندوق الاستثماري صندوق النعيم العقاري والإحاطة علماً بكل الشروط والأحكام المنظمة للاستثمار، نعلمكم برغبتنا في المشاركة، ونلتزم التزاماً كاملاً بشروط وأحكام الصندوق.'
OPEN_EN='Having reviewed the terms & conditions of the Al Naeem Real Estate investment fund thereof, I hereby request to participate in your fund and assure I will abide by the terms and conditions of the investment fund.'
DECL_AR='بهذا أقر بصحة البيانات أعلاه وأن شركة إتقان كابيتال سوف تتعامل معها بسرية تامة، كما أقبل جميع المخاطر المصاحبة للاشتراك بصندوق الاستثمار بما فيها الخسائر الناتجة عن تحويل العملات والتعامل و / أو المتاجرة في الأوراق المالية والمعاملات الأخرى، كما أقر باستلامي نسخة من الشروط والأحكام الخاصة بصندوق الاستثمار.'
DECL_EN='I/We certify that the information above is correct and is provided in the condition that it remains confidential, as well as acknowledging a full understanding of the risks of participating in the investment fund including foreign exchange losses involved in investing, dealing and/or trading in securities and other financial instruments, in addition of receiving the terms and conditions of the investment fund.'

def run(p,txt,size=10,ar=False,bold=False,color=INK):
 r=p.add_run(txt);r.font.name=FONT if ar else 'Arial';r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color);r.bold=bold
 rp=r._r.get_or_add_rPr();rp.rFonts.set(qn('w:cs'),FONT)
 if ar:
  el=OxmlElement('w:rtl');rp.append(el)
  el=OxmlElement('w:szCs');el.set(qn('w:val'),str(int(size*2)));rp.append(el)
  if bold:rp.append(OxmlElement('w:bCs'))
 return r

def para(p,txt='',size=10,ar=False,bold=False,color=INK,leading=None):
 p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(leading or size*1.4)
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if ar else WD_ALIGN_PARAGRAPH.LEFT
 if ar:
  p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
  p._p.get_or_add_pPr().get_or_add_jc().set(qn('w:val'),'start')
 run(p,txt,size,ar,bold,color);return p

def font_embed(path):
 font=bytearray(Path('/Users/ahmaddalao/Library/Fonts/BahijTheSansArabic-Plain.ttf').read_bytes());key=uuid.uuid4();mask=bytes.fromhex(key.hex)[::-1]
 for i in range(32):font[i]^=mask[i%16]
 with zipfile.ZipFile(path) as z:files={n:z.read(n) for n in z.namelist()}
 w='http://schemas.openxmlformats.org/wordprocessingml/2006/main';r='http://schemas.openxmlformats.org/officeDocument/2006/relationships';ns='http://schemas.openxmlformats.org/package/2006/relationships'
 fonts=etree.fromstring(files['word/fontTable.xml']);entry=next((n for n in fonts if n.get('{'+w+'}name')==FONT),None)
 if entry is None:entry=etree.SubElement(fonts,'{'+w+'}font',{'{'+w+'}name':FONT})
 etree.SubElement(entry,'{'+w+'}embedRegular',{'{'+r+'}id':'rIdBrandFont','{'+w+'}fontKey':'{'+str(key).upper()+'}','{'+w+'}subsetted':'false'})
 rels=etree.Element('{'+ns+'}Relationships',nsmap={None:ns});etree.SubElement(rels,'{'+ns+'}Relationship',Id='rIdBrandFont',Type=r+'/font',Target='fonts/brand.odttf')
 files['word/fontTable.xml']=etree.tostring(fonts);files['word/_rels/fontTable.xml.rels']=etree.tostring(rels);files['word/fonts/brand.odttf']=bytes(font)
 ct=etree.fromstring(files['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont');files['[Content_Types].xml']=etree.tostring(ct)
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,data in files.items():z.writestr(n,data)

LABELS={
 'client_account':('Client / Account No. (fund manager use)','رقم العميل / الحساب (لاستخدام مدير الصندوق)'),
 'full_name':('Customer full name','اسم العميل كاملاً'), 'company_name':('Full legal company name','الاسم القانوني الكامل للشركة'),
 'nationality':('Nationality','الجنسية'),'inc_country':('Country of incorporation','دولة التأسيس'),
 'id_type_label':('ID type','نوع الهوية'),'id_number':('ID number','رقم الهوية'),'id_other':('Other identity document','بيان نوع الهوية الأخرى'),
 'company_id_type_label':('Registration type','نوع تسجيل الشركة'),'company_id_number':('Registration / licence number','رقم السجل / الترخيص'),
 'phone':('Telephone','الهاتف'),'mobile':('Mobile','الجوال'),
 'short_address':('Short address','العنوان المختصر'),'building':('Building number','رقم المبنى'),'street':('Street name','اسم الشارع'),
 'additional':('Additional number','الرقم الفرعي'),'district':('District','اسم الحي'),'postal':('Postal code','الرمز البريدي'),'city':('City','المدينة'),'country':('Country','البلد'),'email':('Email','البريد الإلكتروني'),
 'subscription_type_label':('Subscription type','نوع الاشتراك'),'payment_method_label':('Payment method','طريقة الدفع'),
 'fund_name':('Investment fund name','اسم صندوق الاستثمار'),'currency':('Currency','العملة'),'units':('Number of units','عدد الوحدات'),'unit_price':('Unit price (SAR)','سعر الوحدة (ريال سعودي)'),
 'amount_subscribed':('Investment amount (SAR)','مبلغ الاستثمار (ريال سعودي)'), 'subscription_fee':('Subscription fee (2% of investment)','رسوم الاشتراك (2% من مبلغ الاستثمار)'),
 'total_amount':('Total amount (SAR)','المبلغ الإجمالي (ريال سعودي)'), 'total_words':('Total amount in words','المبلغ الإجمالي كتابة'),
 'applicant_name':('Applicant name','اسم مقدم الطلب'),'date':('Application date','التاريخ'),'signature':('Applicant signature','توقيع مقدم الطلب'),
 'auth_name':('Authorized signatory','المفوض بالتوقيع'), 'auth_id':('Signatory ID number','رقم هوية المفوض')}

def build(corporate,probe):
 doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=Pt(37);sec.bottom_margin=Pt(61);sec.left_margin=sec.right_margin=Pt(48);sec.header_distance=Pt(12);sec.footer_distance=Pt(12)
 doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
 doc.styles['Normal'].paragraph_format.space_after=Pt(0)
 doc.core_properties.title='طلب الإشتراك في صندوق النعيم العقاري '+('(للشركات)' if corporate else '(للأفراد)');doc.core_properties.author='';doc.core_properties.last_modified_by=''
 f=sec.footer.paragraphs[0];para(f,'إتقان كابيتال | www.itqancapital.com | +966 12 263 8787',8,True,color=MUTED,leading=11)
 p=sec.footer.add_paragraph();para(p,'الأصل: العمليات والحفظ • نسخة العميل   |   Original: operations & custody • Client copy',7.5,True,color=MUTED,leading=10)
 p=sec.footer.add_paragraph();para(p,'7855 أحمد العطاس، حي الزهراء، وحدة 2563، جدة 23425-2753، المملكة العربية السعودية',6.5,True,color=MUTED,leading=9)
 p=sec.footer.add_paragraph();para(p,'شركة مساهمة سعودية مقفلة • رأس المال: 56,042,030 ريال سعودي • سجل تجاري: 4030167335 • ترخيص هيئة السوق المالية: 37-07058',6.5,True,color=MUTED,leading=9)
 p=sec.footer.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
 markers={}
 def spacer(h=6):
  p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(h);p.paragraph_format.space_after=Pt(0);run(p,' ',1)
 def brand(page):
  p=para(doc.add_paragraph(),'إتقان كابيتال',14,True,True,PURPLE,19)
  para(doc.add_paragraph(),'ITQAN CAPITAL',9,False,True,PURPLE,13)
  if page==1:
   para(doc.add_paragraph(),doc.core_properties.title,17,True,True,INK,25)
   para(doc.add_paragraph(),'Al Naeem Real Estate Fund • Subscription Application • '+('Company' if corporate else 'Individual'),10,False,False,MUTED,17)
  else:
   para(doc.add_paragraph(),'صندوق النعيم العقاري | '+('طلب اشتراك للشركات' if corporate else 'طلب اشتراك للأفراد'),14,True,True,INK,22)
   para(doc.add_paragraph(),'Al Naeem Real Estate Fund • '+('Company' if corporate else 'Individual')+' Subscription',9,False,False,MUTED,14)
  spacer(7)
 def heading(en,ar):
  p=para(doc.add_paragraph(),ar,12,True,True,PURPLE,18);p.paragraph_format.space_before=Pt(6)
  para(doc.add_paragraph(),en,8.5,False,True,MUTED,13);spacer(3)
 def row(ids,height=47,widths=None,fixed=None):
  widths=widths or [WIDTH/len(ids)]*len(ids);table=doc.add_table(rows=1,cols=len(ids));table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','bottom','left','right','insideV','insideH']:
   n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'3');n.set(qn('w:color'),BORDER);borders.append(n)
  table._tbl.tblPr.append(borders);table.rows[0].height=Pt(height);table.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
  cant=OxmlElement('w:cantSplit');table.rows[0]._tr.get_or_add_trPr().append(cant)
  for i,(key,w) in enumerate(zip(ids,widths)):
   cell=table.cell(0,i);cell.width=Pt(w);table.columns[i].width=Pt(w)
   margins=OxmlElement('w:tcMar')
   for side,n in [('top',35),('bottom',25),('left',100),('right',100)]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),str(n));el.set(qn('w:type'),'dxa');margins.append(el)
   cell._tc.get_or_add_tcPr().append(margins)
   en,ar=LABELS[key];para(cell.paragraphs[0],ar,9.5,True,color=INK,leading=13)
   para(cell.add_paragraph(),en,7.5,False,color=MUTED,leading=11)
   p=cell.add_paragraph();p.paragraph_format.line_spacing=Pt(height-28);p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
   token='S'+str(len(markers)+1).zfill(3)+'X';markers[token]={'id':key,'width':w-12,'height':height-29}
   if fixed and key in fixed:run(p,fixed[key],11,True,color='1456A0')
   else:run(p,token if probe else ' ',4)
  return table
 brand(1)
 para(doc.add_paragraph(),'السادة / شركة إتقان كابيتال',10,True,True,leading=15)
 para(doc.add_paragraph(),'Messrs.: Itqan Capital Co.',8,False,False,MUTED,12)
 para(doc.add_paragraph(),OPEN_AR,9,True,leading=13)
 para(doc.add_paragraph(),OPEN_EN,8,False,leading=10.5)
 heading('Customer details','تفاصيل العميل')
 row(['client_account'],43)
 row(['company_name' if corporate else 'full_name'],47)
 if corporate:
  row(['inc_country','company_id_type_label'],47)
  row(['company_id_number','auth_id'],47)
  row(['auth_name'],47)
 else:
  row(['nationality','id_type_label'],47)
  row(['id_number','id_other'],47)
 row(['phone','mobile'],47)
 heading('Correspondence address (National Address)','عنوان المراسلة (العنوان الوطني)')
 row(['short_address','building','additional'],47)
 row(['street','district'],47)
 row(['city','postal','country'],47)
 row(['email'],47)
 doc.add_page_break();brand(2)
 heading('Subscription details','تفاصيل الاشتراك')
 row(['subscription_type_label','payment_method_label'],47)
 row(['fund_name','currency'],47,widths=[WIDTH*.7,WIDTH*.3],fixed={'fund_name':RULES['fundName'],'currency':RULES['currency']})
 row(['units','unit_price'],47,fixed={'unit_price':format(RULES['unitPrice'],',')})
 row(['amount_subscribed','subscription_fee'],47)
 row(['total_amount'],47)
 row(['total_words'],53)
 spacer(6)
 para(doc.add_paragraph(),DECL_AR,9,True,leading=13)
 para(doc.add_paragraph(),DECL_EN,8,False,leading=10.5)
 heading('Applicant' if not corporate else 'Authorized applicant','مقدم الطلب' if not corporate else 'مقدم الطلب المفوض')
 row(['applicant_name','date'],47,widths=[WIDTH*.7,WIDTH*.3])
 row(['signature'],66)
 spacer(6)
 # The internal approval block stays on paper; clients do not complete it online.
 para(doc.add_paragraph(),'لاستخدام مدير الصندوق فقط | For fund manager use only',9,True,True,PURPLE,13)
 para(doc.add_paragraph(),'التوقيع مطابق ☐    الفرع: __________________    التاريخ: __________________',8.5,True,leading=15)
 para(doc.add_paragraph(),'مدير الحساب وتوقيعه: __________________    مدخل الطلب وتوقيعه: __________________',8.5,True,leading=15)
 para(doc.add_paragraph(),'المراجع والمعتمد وتوقيعه: __________________',8.5,True,leading=15)
 return doc,markers

for corporate in [False,True]:
 name='subscription-company' if corporate else 'subscription-individual';base=ROOT/'tmp/subscription'/name;base.mkdir(parents=True,exist_ok=True)
 maps={}
 for probe in [True,False]:
  doc,markers=build(corporate,probe);target=base/'probe.docx' if probe else ROOT/'output/documents'/f'{name}.docx';doc.save(target);font_embed(target)
  out=base/('probe' if probe else 'final')
  subprocess.run([str(RUNTIME/'python/bin/python3'),str(RENDER),str(target),'--output_dir',str(out),'--emit_pdf'],check=True)
  pdf=out/(target.stem+'.pdf')
  with pdfplumber.open(pdf) as rendered:
   assert len(rendered.pages)==2,(name,len(rendered.pages))
   if probe:
    for page_number,page in enumerate(rendered.pages,1):
     for token,info in markers.items():
      for match in page.search(token,regex=False):
       bottom=min(edge['top'] for edge in page.edges if edge['orientation']=='h' and edge['x0']<=match['x0']<=edge['x1'] and edge['top']>=match['bottom'])
       maps[info['id']]={'page':page_number,'rect':[round(match['x0'],2),round(bottom-info['height']-4,2),round(info['width'],2),round(info['height'],2)]}
   else:
    assert all(not page.search(r'S\d{3}X') for page in rendered.pages)
  if not probe:
   shutil.copyfile(pdf,ROOT/'output/pdf'/f'{name}.pdf')
   shutil.copyfile(pdf,ROOT/'public/pdfs'/f'{name}.pdf')
 (ROOT/'src/subscription'/f'{name}-layout.json').write_text(json.dumps(maps,ensure_ascii=False,indent=2)+'\n')
 print(name,len(maps),'mapped fields')
