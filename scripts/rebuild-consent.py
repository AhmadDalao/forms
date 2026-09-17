"""Recreate the supplied consent page as editable Arabic text, without page 38."""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/documents/al-naeem-terms-consent.docx'
BODY = 'لقد قمت / قمنا بقراءة الشروط والأحكام والملاحق الخاصة بالصندوق وفهم ما جاء فيها والموافقة عليها، كما جرى الحصول على نسخة منها بعد التوقيع عليها. وإثباتاً لما تقدم، قام المستثمر بالتوقيع على هذه الشروط والأحكام الخاصة بالصندوق في التاريخ والسنة المذكورين أدناه.'


def arabic(paragraph, text='', size=14, bold=False, underline=False):
    paragraph._p.get_or_add_pPr().get_or_add_jc().set(qn('w:val'), 'start')
    paragraph.paragraph_format.line_spacing = Pt(24)
    paragraph.paragraph_format.space_after = Pt(0)
    bidi = OxmlElement('w:bidi')
    paragraph._p.get_or_add_pPr().append(bidi)
    if text:
        run = paragraph.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = bold
        run.underline = underline
        pr = run._r.get_or_add_rPr()
        pr.rFonts.set(qn('w:cs'), 'Arial')
        for tag, value in [('rtl', None), ('szCs', str(size * 2)), ('lang', 'ar-SA')]:
            node = OxmlElement('w:' + tag)
            if value:
                node.set(qn('w:bidi' if tag == 'lang' else 'w:val'), value)
            pr.append(node)
        if bold:
            pr.append(OxmlElement('w:bCs'))
    return paragraph


def plain_table(table):
    table.autofit = False
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:val'), 'nil')
        borders.append(edge)
    table._tbl.tblPr.append(borders)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(24), Mm(22)
section.left_margin, section.right_margin = Mm(22), Mm(22)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Arial', Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'إقرار من مالكي الوحدات - صندوق النعيم العقاري'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''

p = arabic(doc.add_paragraph(), 'إتقان كابيتال | صندوق النعيم العقاري', 12)
p.paragraph_format.space_after = Pt(26)
p = arabic(doc.add_paragraph(style='Title'), 'إقرار من مالكي الوحدات:', 15, True, True)
p.paragraph_format.space_after = Pt(5)
p = arabic(doc.add_paragraph(), BODY)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after = Pt(42)
p = arabic(doc.add_paragraph(), 'من قبل المستثمر', 15, True, True)
p.paragraph_format.space_after = Pt(12)

table = doc.add_table(rows=3, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.RIGHT
plain_table(table)
table.columns[0].width, table.columns[1].width = Mm(84), Mm(20)
for row, label in zip(table.rows, ['الاسم:', 'التوقيع:', 'التاريخ:']):
    row.height, row.height_rule = Mm(13), WD_ROW_HEIGHT_RULE.EXACTLY
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row.cells[0].width, row.cells[1].width = Mm(84), Mm(20)
    p = arabic(row.cells[0].paragraphs[0])
    p.paragraph_format.line_spacing = Pt(20)
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for key, value in [('val', 'single'), ('sz', '4'), ('color', '555555'), ('space', '1')]:
        bottom.set(qn('w:' + key), value)
    borders.append(bottom)
    p._p.get_or_add_pPr().append(borders)
    arabic(row.cells[1].paragraphs[0], label, 14)

p = arabic(doc.add_paragraph())
p.paragraph_format.space_after = Pt(87)
p.paragraph_format.line_spacing = Pt(1)
officials = doc.add_table(rows=1, cols=2)
plain_table(officials)
officials.columns[0].width = officials.columns[1].width = Mm(83)
for cell, lines in zip(officials.rows[0].cells, [
    ['العضو المنتدب والرئيس التنفيذي', 'د. محمد بسام هاشم السيد'],
    ['مسؤول المطابقة والالتزام', 'أسامه فايز المالكي'],
]):
    cell.width = Mm(83)
    for i, line in enumerate(lines):
        p = arabic(cell.paragraphs[0] if i == 0 else cell.add_paragraph(), line, 14, True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
